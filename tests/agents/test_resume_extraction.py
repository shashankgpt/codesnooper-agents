import io
import subprocess
import threading
from typing import Any, Literal
from unittest.mock import Mock

import httpx
import pytest
from docx import Document

from app.agents import resume_parser_agent as parser


@pytest.mark.parametrize("document_type", ["doc", "docx"])
async def test_word_extraction_runs_outside_event_loop(
    monkeypatch: pytest.MonkeyPatch, document_type: Literal["doc", "docx"]
) -> None:
    loop_thread = threading.get_ident()
    worker_threads: list[int] = []

    def parse_in_thread(content: bytes, kind: str) -> str:
        worker_threads.append(threading.get_ident())
        assert content == b"document" and kind == document_type
        return "Resume text"

    monkeypatch.setattr(parser, "_extract_word_text", parse_in_thread)
    assert await parser.extract_word_text(b"document", document_type) == "Resume text"
    assert worker_threads and worker_threads[0] != loop_thread


async def test_docx_extracts_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Resume summary")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Python"
    row.cells[1].text = "FastAPI"
    content = io.BytesIO()
    document.save(content)
    assert await parser.extract_word_text(content.getvalue(), "docx") == (
        "Resume summary\nPython | FastAPI"
    )


@pytest.mark.parametrize(
    "failure",
    [FileNotFoundError("antiword"), subprocess.TimeoutExpired("antiword", 20)],
)
async def test_antiword_errors_are_friendly(
    monkeypatch: pytest.MonkeyPatch, failure: Exception
) -> None:
    monkeypatch.setattr(parser.subprocess, "run", Mock(side_effect=failure))
    with pytest.raises(ValueError, match="Please try saving it as DOCX or PDF"):
        await parser.extract_word_text(b"document", "doc")


@pytest.mark.parametrize("returncode, stdout", [(1, b""), (0, b"")])
async def test_unreadable_doc_is_rejected(
    monkeypatch: pytest.MonkeyPatch, returncode: int, stdout: bytes
) -> None:
    monkeypatch.setattr(
        parser.subprocess,
        "run",
        Mock(return_value=subprocess.CompletedProcess("antiword", returncode, stdout=stdout)),
    )
    with pytest.raises(ValueError, match="Could not"):
        await parser.extract_word_text(b"document", "doc")


@pytest.mark.parametrize("from_bytes", [True, False])
async def test_pdf_extraction_runs_outside_event_loop(
    monkeypatch: pytest.MonkeyPatch, from_bytes: bool
) -> None:
    loop_thread = threading.get_ident()
    worker_threads: list[int] = []

    def parse_in_thread(content: bytes) -> str:
        worker_threads.append(threading.get_ident())
        assert content == b"PDF bytes"
        return "PDF text"

    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=b"PDF bytes"))
    async_client = httpx.AsyncClient

    def make_client(**kwargs: Any) -> httpx.AsyncClient:
        return async_client(transport=transport, **kwargs)

    monkeypatch.setattr(parser, "extract_pdf_text", parse_in_thread)
    monkeypatch.setattr(parser.httpx, "AsyncClient", make_client)
    assert await parser.fetch_source_text(
        "https://example.com/resume.pdf", source_bytes=b"PDF bytes" if from_bytes else None
    ) == ("PDF text", "pdf")
    assert worker_threads and worker_threads[0] != loop_thread
